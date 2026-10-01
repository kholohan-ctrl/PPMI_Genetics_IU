#!/usr/bin/env python3
"""
mergefastqXlsx.py

Read an XLSX file with columns:
  - gp2id
  - path   (R1 FASTQ paths only)

For each gp2id:
  - group all R1 paths from the XLSX
  - infer the matching R2 paths by replacing _R1_001.fastq.gz with _R2_001.fastq.gz
  - merge R1 and R2 separately

GCP actions are done with gsutil shell commands.
"""

from __future__ import annotations

import argparse
import logging
import re
import shutil
import subprocess
import sys
import tempfile
from collections import defaultdict
from pathlib import Path
from typing import Dict, List, Tuple

from openpyxl import load_workbook


LOG = logging.getLogger("mergefastqXlsx")

R1_RE = re.compile(r"_R1_001\.fastq\.gz$")
R2_RE = re.compile(r"_R2_001\.fastq\.gz$")


def run_cmd(cmd: List[str]) -> None:
    subprocess.run(cmd, check=True)


def infer_r2_path(r1_path: str) -> str:
    if not R1_RE.search(r1_path):
        raise ValueError(f"Not an R1 FASTQ path: {r1_path}")
    return R1_RE.sub("_R2_001.fastq.gz", r1_path)


def read_xlsx_rows(xlsx_source: str, sheet_name: str | None) -> List[Tuple[str, str]]:
    with tempfile.TemporaryDirectory() as td:
        local_xlsx = Path(td) / "input.xlsx"

        if xlsx_source.startswith("gs://"):
            run_cmd(["gsutil", "cp", xlsx_source, str(local_xlsx)])
        else:
            shutil.copy2(xlsx_source, local_xlsx)

        wb = load_workbook(local_xlsx, read_only=True, data_only=True)

        if sheet_name:
            if sheet_name not in wb.sheetnames:
                raise ValueError(f"Sheet not found: {sheet_name}")
            ws = wb[sheet_name]
        else:
            ws = wb[wb.sheetnames[0]]

        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            raise ValueError("XLSX has no header row")

        header_map: Dict[str, int] = {}
        for i, col in enumerate(header):
            if col is not None:
                header_map[str(col).strip().lower()] = i

        if "gp2id" not in header_map:
            raise ValueError("XLSX is missing required column: gp2id")
        if "path" not in header_map:
            raise ValueError("XLSX is missing required column: path")

        gp2id_idx = header_map["gp2id"]
        path_idx = header_map["path"]

        out: List[Tuple[str, str]] = []
        for row in rows:
            if row is None:
                continue
            gp2id = row[gp2id_idx] if gp2id_idx < len(row) else None
            path = row[path_idx] if path_idx < len(row) else None
            if gp2id is None or path is None:
                continue
            gp2id = str(gp2id).strip()
            path = str(path).strip()
            if gp2id and path:
                out.append((gp2id, path))

        return out


def group_r1_r2(rows: List[Tuple[str, str]]) -> Dict[str, Dict[str, List[str]]]:
    grouped: Dict[str, Dict[str, List[str]]] = defaultdict(lambda: {"R1": [], "R2": []})

    for gp2id, r1_path in rows:
        base = r1_path.rsplit("/", 1)[-1]
        if not R1_RE.search(base):
            LOG.warning("Skipping non-R1 path in XLSX: %s", r1_path)
            continue

        r2_path = infer_r2_path(r1_path)

        grouped[gp2id]["R1"].append(r1_path)
        grouped[gp2id]["R2"].append(r2_path)

    for gp2id in grouped:
        grouped[gp2id]["R1"] = sorted(set(grouped[gp2id]["R1"]))
        grouped[gp2id]["R2"] = sorted(set(grouped[gp2id]["R2"]))

    return grouped


def gs_merge_to_gcs(sources: List[str], dest_uri: str) -> None:
    cat_cmd = ["gsutil", "cat", *sources]
    cp_cmd = ["gsutil", "cp", "-", dest_uri]

    cat_proc = subprocess.Popen(cat_cmd, stdout=subprocess.PIPE)
    assert cat_proc.stdout is not None

    try:
        cp_proc = subprocess.Popen(cp_cmd, stdin=cat_proc.stdout)
        cat_proc.stdout.close()
        cp_rc = cp_proc.wait()
        cat_rc = cat_proc.wait()

        if cat_rc != 0:
            raise subprocess.CalledProcessError(cat_rc, cat_cmd)
        if cp_rc != 0:
            raise subprocess.CalledProcessError(cp_rc, cp_cmd)
    finally:
        if cat_proc.poll() is None:
            cat_proc.kill()


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Merge FASTQs listed in an XLSX file, one gp2id at a time."
    )
    parser.add_argument("--input-xlsx", required=True, help="Local path or gs:// URI to XLSX")
    parser.add_argument("--output-prefix", required=True, help="gs:// output prefix")
    parser.add_argument("--sheet-name", default="", help="Optional worksheet name")
    parser.add_argument("--dry-run", action="store_true", help="Print actions only")
    parser.add_argument("--log-level", default="INFO", help="DEBUG, INFO, WARNING, ERROR")
    parser.add_argument("--check-r2-exists", action="store_true",
                        help="Verify inferred R2 files exist in GCS before merging")

    args = parser.parse_args()

    if not args.output_prefix.startswith("gs://"):
        raise SystemExit("--output-prefix must be a gs:// URI")

    logging.basicConfig(
        level=getattr(logging, args.log_level.upper(), logging.INFO),
        format="%(message)s",
    )

    LOG.info("==================================================")
    LOG.info("Starting mergefastqXlsx.py")
    LOG.info("==================================================")
    LOG.info("Input XLSX : %s", args.input_xlsx)
    LOG.info("Output dir : %s", args.output_prefix)
    LOG.info("Sheet name : %s", args.sheet_name or "<first sheet>")
    LOG.info("Dry run    : %s", args.dry_run)
    LOG.info("")

    rows = read_xlsx_rows(args.input_xlsx, args.sheet_name or None)
    LOG.info("Parsed %d data rows from XLSX.", len(rows))
    LOG.info("")

    grouped = group_r1_r2(rows)
    samples = sorted(grouped.keys())

    LOG.info("Found %d samples.", len(samples))
    LOG.info("")

    for sample_id in samples:
        r1_files = grouped[sample_id]["R1"]
        r2_files = grouped[sample_id]["R2"]

        LOG.info("==================================================")
        LOG.info("Processing: %s", sample_id)
        LOG.info("==================================================")

        if not r1_files:
            LOG.warning("No R1 files found")
            LOG.info("")
            continue

        if len(r1_files) != len(r2_files):
            LOG.error("R1/R2 count mismatch")
            LOG.error("R1 count: %d", len(r1_files))
            LOG.error("R2 count: %d", len(r2_files))
            LOG.info("")
            continue

        out_r1 = f"{args.output_prefix.rstrip('/')}/{sample_id}_merged_R1.fastq.gz"
        out_r2 = f"{args.output_prefix.rstrip('/')}/{sample_id}_merged_R2.fastq.gz"

        LOG.info("R1 files:")
        for p in r1_files:
            LOG.info("  %s", p)
        LOG.info("")

        LOG.info("R2 files:")
        for p in r2_files:
            LOG.info("  %s", p)
        LOG.info("")

        if args.check_r2_exists:
            missing = []
            for p in r2_files:
                if not p.startswith("gs://"):
                    missing.append(p)
                    continue
                # shell check via gsutil stat
                try:
                    subprocess.run(["gsutil", "stat", p], check=True,
                                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
                except subprocess.CalledProcessError:
                    missing.append(p)
            if missing:
                LOG.error("Missing inferred R2 files:")
                for p in missing:
                    LOG.error("  %s", p)
                LOG.info("")
                continue

        if args.dry_run:
            LOG.info("[DRY RUN] gsutil cat %s | gsutil cp - %s", " ".join(r1_files), out_r1)
            LOG.info("[DRY RUN] gsutil cat %s | gsutil cp - %s", " ".join(r2_files), out_r2)
            LOG.info("")
            continue

        LOG.info("Merging R1 -> %s", out_r1)
        gs_merge_to_gcs(r1_files, out_r1)

        LOG.info("Merging R2 -> %s", out_r2)
        gs_merge_to_gcs(r2_files, out_r2)

        LOG.info("Finished sample: %s", sample_id)
        LOG.info("")

    LOG.info("==================================================")
    LOG.info("All samples processed.")
    LOG.info("==================================================")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
