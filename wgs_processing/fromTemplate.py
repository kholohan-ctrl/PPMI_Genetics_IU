import csv
from openpyxl import load_workbook

batchname = '20260817n'
batchdir = 'ngdpr_20260817'
dir = 'C:/Users/jbmchls/Indiana University/Nudelman, Kelly Nicole Holohan - Jessie_Projects/PPMI/WGS_processing/'
csv_in = f"{dir}mz_scripts/multiinputs_template.csv"
xlsx_in = f"{dir}{batchdir}/r1_files{batchname}.xlsx"
# xlsx_in = f"{dir}r1_files1.xlsx_multiple-5_2parts.xlsx"
csv_out = f"{dir}{batchdir}/multiinputs-{batchname}-path.csv"

# Read the template CSV
with open(csv_in, newline="", encoding="utf-8") as f:
    reader = csv.DictReader(f)
    fieldnames = reader.fieldnames
    template_rows = list(reader)

if not fieldnames:
    raise ValueError("CSV has no header row.")

if not template_rows:
    raise ValueError("CSV has no data row to use as the template.")

template = template_rows[0]

# Read the Excel file and get all values from the column named 'path'
wb = load_workbook(xlsx_in, read_only=True, data_only=True)
ws = wb.active

header = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
if "path" not in header:
    raise ValueError("Excel file does not contain a column named 'path'.")

path_col_idx = header.index("path") + 1

paths = []
for row in ws.iter_rows(min_row=2):
    value = row[path_col_idx - 1].value
    if value is not None and str(value).strip() != "":
        paths.append(str(value).strip())

# Build output rows
output_rows = []
for p in paths:
    row = dict(template)  # copy all other columns from first data row
    row["fastq_to_cram.fastq_R1"] = p
    row["fastq_to_cram.fastq_R2"] = p.replace("R1", "R2")
    output_rows.append(row)

# Write the populated CSV
with open(csv_out, "w", newline="", encoding="utf-8") as f:
    writer = csv.DictWriter(f, fieldnames=fieldnames)
    writer.writeheader()
    writer.writerows(output_rows)

print(f"Wrote {len(output_rows)} rows to {csv_out}")